"""Write the summary tables to a formatted .xlsx workbook (local file, offline)."""

from __future__ import annotations

import os

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .aggregate import Summaries

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=12)


def _write_sheet(writer: pd.ExcelWriter, name: str, df: pd.DataFrame) -> None:
    if df is None or df.empty:
        df = pd.DataFrame({"(no data)": []})
    df.to_excel(writer, sheet_name=name, index=False)
    ws = writer.sheets[name]

    # style header row
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # auto-ish column widths
    for col_idx, col in enumerate(df.columns, start=1):
        max_len = max(
            [len(str(col))]
            + [len(str(v)) for v in df.iloc[:, col_idx - 1].tolist()]
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)

    ws.freeze_panes = "A2"


def write_workbook(summaries: Summaries, output_path: str) -> str:
    """Write all sheets to output_path. Returns the path."""
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        if summaries.daily is not None and not summaries.daily.empty:
            _write_sheet(writer, "Daily Summary", summaries.daily)
        _write_sheet(writer, "Summary", summaries.summary)
        _write_sheet(writer, "Counts by Type", summaries.counts_by_type)
        _write_sheet(writer, "TAT by Type", summaries.tat_by_type)
        _write_sheet(writer, "TAT by Type & Shift", summaries.tat_by_type_shift)
        _write_sheet(writer, "Detail", summaries.detail)
    return output_path


def append_to_master(daily_rows: pd.DataFrame, master_path: str) -> str:
    """Append the Daily Summary row(s) to a running master workbook.

    A report can now yield several rows (one per date), and all are appended.
    Creates the file (with a styled header) if it doesn't exist. If it does,
    rows are aligned to the existing header by column name — any columns the
    master doesn't have yet are added on the end — so the layout can evolve
    without breaking older rows. Returns the master path.
    """
    if daily_rows is None or daily_rows.empty:
        return master_path

    dict_rows = [r.to_dict() for _, r in daily_rows.iterrows()]

    if not os.path.exists(master_path):
        os.makedirs(os.path.dirname(master_path) or ".", exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Master"
        headers = list(daily_rows.columns)
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in dict_rows:
            ws.append([row.get(h) for h in headers])
        ws.freeze_panes = "A2"
        wb.save(master_path)
        return master_path

    wb = openpyxl.load_workbook(master_path)
    ws = wb["Master"] if "Master" in wb.sheetnames else wb.active
    headers = [c.value for c in ws[1]]

    # add any brand-new columns to the header
    for col in daily_rows.columns:
        if col not in headers:
            headers.append(col)
            cell = ws.cell(row=1, column=len(headers), value=col)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in dict_rows:
        ws.append([row.get(h) for h in headers])
    wb.save(master_path)
    return master_path
